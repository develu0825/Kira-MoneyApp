"""
Kira CSV 가져오기 — 은행 앱에서 내보낸 거래내역(엑셀/CSV)을 money.db에 넣는다.
CODEF 없이 바로 쓸 수 있는 경로.

사용법:
  카뱅/토스 앱 → 거래내역 → 내보내기(엑셀/CSV) → 이 폴더에 저장
  python import_csv.py 내파일.csv
  python import_csv.py 내파일.xlsx --bank 카카오뱅크
  python import_csv.py *.csv            (여러 개 한 번에)

그 다음:  python publish.py   (암호화 → 폰 앱으로)

지원: 대부분의 한국 은행 CSV/엑셀. 컬럼명을 자동 인식하고,
      못 찾으면 어떤 컬럼이 있는지 출력해주니 그걸 알려주면 맞춰줄 수 있다.
"""

import os
import re
import sys
import csv
import glob
import sqlite3
import datetime as dt
from pathlib import Path

# 윈도우 콘솔(cp949/cp1252)에서 한글·이모지 출력이 깨지지 않도록
try:
    sys.stdout.reconfigure(encoding="utf-8")
    sys.stderr.reconfigure(encoding="utf-8")
except Exception:
    pass

HERE = Path(__file__).parent
DB_PATH = HERE / "money.db"

try:
    from dotenv import load_dotenv
    load_dotenv(HERE / ".env")
except ImportError:
    pass

try:
    from categorize import guess_category
except ImportError:
    def guess_category(_name):
        return None
DOW_KR = ["월", "화", "수", "목", "금", "토", "일"]

# 컬럼 자동 매칭용 키워드 (은행마다 이름이 다름)
COLS = {
    "date":    ["거래일시", "거래일자", "거래일", "일자", "날짜", "승인일자", "date"],
    "time":    ["거래시간", "시간", "거래시각", "time"],
    "out":     ["출금", "출금액", "출금금액", "보낸금액", "지출", "결제금액", "withdraw"],
    "in":      ["입금", "입금액", "입금금액", "받은금액", "수입", "deposit"],
    "amount":  ["거래금액", "금액", "amount"],      # 부호로 입출금 구분하는 형식
    "balance": ["잔액", "거래후잔액", "잔고", "balance"],
    "desc":    ["적요", "내용", "거래내용", "받는분", "보내는분", "상대방", "가맹점명",
                "거래처", "메모", "기재내용", "description", "메모내용", "거래기록사항"],
}


def _norm(s: str) -> str:
    return re.sub(r"[\s_()\[\]./-]", "", (s or "").strip().lower())


def _find_col(headers, keys):
    nheaders = {_norm(h): h for h in headers}
    for key in keys:
        nk = _norm(key)
        for nh, orig in nheaders.items():
            if nk and nk in nh:
                return orig
    return None


def _to_int(v) -> int:
    if v is None:
        return 0
    s = re.sub(r"[^\d\-]", "", str(v))
    return int(s) if s and s not in ("-", "") else 0


def _open_excel_bytes(path: str):
    """엑셀 파일 바이트 반환. 암호화(카뱅 등)면 암호로 복호화. 구형 .xls면 그대로."""
    import io
    with open(path, "rb") as f:
        head = f.read(8)
    # OLE2(d0cf11e0) = 구형 xls 또는 암호화 xlsx
    if head[:4] == b"\xd0\xcf\x11\xe0":
        try:
            import msoffcrypto
        except ImportError:
            sys.exit("암호화 엑셀을 열려면:  pip install msoffcrypto-tool")
        with open(path, "rb") as f:
            office = msoffcrypto.OfficeFile(f)
            if not office.is_encrypted():
                return None  # 암호화 아닌 구형 xls → 호출측에서 xlrd 처리
            pw = (os.environ.get("XLSX_PW")
                  or os.environ.get("MY_BIRTH", "")
                  or input("엑셀 암호 (보통 생년월일 6자리): ").strip())
            buf = io.BytesIO()
            try:
                office.load_key(password=pw)
                office.decrypt(buf)
            except Exception:
                sys.exit("❌ 엑셀 암호가 틀렸어요. 카뱅 내보내기 때 설정한 암호(보통 생년월일 6자리)로 다시 시도하세요.\n"
                         "   예:  python import_csv.py 파일.xlsx --pw 900101")
            buf.seek(0)
            return buf
    return b""  # 일반 xlsx (zip)


def _read_rows(path: str):
    """CSV 또는 XLSX에서 (헤더, 행리스트) 읽기. 헤더가 몇 줄 아래 있어도 찾아냄."""
    ext = Path(path).suffix.lower()
    if ext in (".xlsx", ".xls"):
        from openpyxl import load_workbook
        src = _open_excel_bytes(path)
        if src is None:
            # 진짜 구형 .xls (암호화 아님)
            import xlrd
            wb_x = xlrd.open_workbook(path)
            sh = wb_x.sheet_by_index(0)
            rows = [sh.row_values(i) for i in range(sh.nrows)]
            return _rows_after_header(rows)
        wb = load_workbook(src if src else path, read_only=True, data_only=True)
        ws = wb.active
        rows = [[("" if c is None else c) for c in r] for r in ws.iter_rows(values_only=True)]
    else:
        raw = None
        for enc in ("utf-8-sig", "cp949", "euc-kr", "utf-8"):
            try:
                with open(path, encoding=enc, newline="") as f:
                    raw = list(csv.reader(f))
                break
            except (UnicodeDecodeError, LookupError):
                continue
        if raw is None:
            sys.exit(f"파일 인코딩을 못 읽었어요: {path}")
        rows = raw

    return _rows_after_header(rows)


def _rows_after_header(rows):
    """헤더 행(COLS 키워드가 가장 많은 줄)을 찾아 (헤더, 데이터행) 반환."""
    header_idx = 0
    best = 0
    for i, r in enumerate(rows[:15]):
        cells = [str(c) for c in r]
        hits = sum(1 for keys in COLS.values() if _find_col(cells, keys))
        if hits > best:
            best, header_idx = hits, i
    headers = [str(c).strip() for c in rows[header_idx]]
    data = [r for r in rows[header_idx + 1:] if any(str(c).strip() for c in r)]
    return headers, data


def _parse_dt(date_s, time_s):
    s = (str(date_s) or "").strip()
    t = (str(time_s) or "").strip()
    # 날짜+시간이 한 칸에 붙어온 경우
    m = re.search(r"(\d{4})[.\-/]?(\d{1,2})[.\-/]?(\d{1,2})\D*(\d{1,2})?:?(\d{2})?", s)
    if not m:
        return None
    y, mo, d = int(m.group(1)), int(m.group(2)), int(m.group(3))
    hh = int(m.group(4) or 0)
    mm = int(m.group(5) or 0)
    if t:
        tm = re.search(r"(\d{1,2}):(\d{2})", t)
        if tm:
            hh, mm = int(tm.group(1)), int(tm.group(2))
    try:
        return dt.datetime(y, mo, d, hh, mm)
    except ValueError:
        return None


def init_db(conn):
    conn.executescript(
        """
        CREATE TABLE IF NOT EXISTS transactions (
            tx_key TEXT PRIMARY KEY, organization TEXT, account TEXT, datetime TEXT,
            in_amount INTEGER DEFAULT 0, out_amount INTEGER DEFAULT 0, balance INTEGER,
            counterparty TEXT, memo_bank TEXT,
            excluded INTEGER DEFAULT 0, exclude_reason TEXT, category TEXT,
            is_self_transfer INTEGER DEFAULT 0, user_memo TEXT, raw_json TEXT, created_at TEXT NOT NULL
        );
        CREATE INDEX IF NOT EXISTS idx_tx_datetime ON transactions(datetime);
        """
    )
    conn.commit()


def import_file(conn, path, bank):
    headers, rows = _read_rows(path)
    c_date = _find_col(headers, COLS["date"])
    c_time = _find_col(headers, COLS["time"])
    c_out = _find_col(headers, COLS["out"])
    c_in = _find_col(headers, COLS["in"])
    c_amt = _find_col(headers, COLS["amount"])
    c_bal = _find_col(headers, COLS["balance"])
    c_desc = _find_col(headers, COLS["desc"])

    if not c_date or not (c_out or c_in or c_amt):
        print(f"\n⚠️ {Path(path).name}: 컬럼 자동 인식 실패.")
        print(f"   이 파일의 컬럼들: {headers}")
        print("   → 이 목록을 보여주면 매핑을 맞춰줄게.")
        return 0

    idx = {h: i for i, h in enumerate(headers)}
    now = dt.datetime.now().isoformat()
    n = 0
    for i, r in enumerate(rows):
        def cell(col):
            return r[idx[col]] if col and idx.get(col) is not None and idx[col] < len(r) else ""

        when = _parse_dt(cell(c_date), cell(c_time))
        if not when:
            continue
        if c_amt and not (c_out or c_in):
            amt = _to_int(cell(c_amt))
            out_amt, in_amt = (abs(amt), 0) if amt < 0 else (0, amt)
        else:
            out_amt = _to_int(cell(c_out))
            in_amt = _to_int(cell(c_in))
        balance = _to_int(cell(c_bal))
        desc = str(cell(c_desc)).strip()
        iso = when.isoformat()
        tx_key = f"csv|{bank}|{iso}|{in_amt}|{out_amt}|{balance}|{i}"
        category = guess_category(desc) if out_amt > 0 else None   # 지출만 자동분류
        conn.execute(
            """INSERT INTO transactions
               (tx_key, organization, account, datetime, in_amount, out_amount, balance,
                counterparty, memo_bank, category, is_self_transfer, raw_json, created_at)
               VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?)
               ON CONFLICT(tx_key) DO UPDATE SET
                 in_amount=excluded.in_amount, out_amount=excluded.out_amount,
                 balance=excluded.balance, counterparty=excluded.counterparty,
                 category=COALESCE(transactions.category, excluded.category)""",
            (tx_key, bank, "", iso, in_amt, out_amt, balance, desc, "", category, 0, "", now),
        )
        n += 1
    conn.commit()
    return n


def main():
    # 플래그와 그 값을 제외한 나머지 = 파일 경로
    argv = sys.argv[1:]
    args, skip = [], False
    for i, a in enumerate(argv):
        if skip:
            skip = False
            continue
        if a in ("--bank", "--pw"):
            skip = True
            continue
        if a.startswith("--"):
            continue
        args.append(a)
    bank = "은행"
    if "--bank" in sys.argv:
        bank = sys.argv[sys.argv.index("--bank") + 1]
    if "--pw" in sys.argv:
        os.environ["XLSX_PW"] = sys.argv[sys.argv.index("--pw") + 1]
    if not args:
        sys.exit("사용법: python import_csv.py 거래내역.csv [--bank 카카오뱅크]")

    files = []
    for a in args:
        files.extend(glob.glob(a))
    if not files:
        sys.exit(f"파일을 못 찾았어요: {args}")

    conn = sqlite3.connect(DB_PATH)
    init_db(conn)
    total = 0
    for f in files:
        n = import_file(conn, f, bank)
        if n:
            print(f"✅ {Path(f).name}: {n}건 가져옴")
            total += n
    conn.close()
    if total:
        print(f"\n총 {total}건 → money.db 저장 완료")
        print("다음:  python publish.py   (암호화해서 폰 앱으로)")


if __name__ == "__main__":
    main()
